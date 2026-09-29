import asyncio
import copy
import hashlib
import json
import csv
import math
import os
import re
import tempfile
import time
from pathlib import Path
from datetime import datetime

# ============================================================
# VERSION / CONFIG
# ============================================================

V31_VERSION = "3.1.0"

V31_ROOT = Path(
    os.environ.get(
        "STRUCTURAL_PROJECT_DATA_DIR",
        "./project_data"
    )
)

V31_ROOT.mkdir(
    parents=True,
    exist_ok=True
)

V31_PROJECT_LIMIT = 200
V31_MEMBER_LIMIT = 5000
V31_FILE_LIMIT_MB = 20
V31_CALLBACK_GUARD_SECONDS = 1.25
V31_CALC_CACHE_LIMIT = 256
V31_EXPORT_TTL_SECONDS = 1800
V31_STOCK_LENGTH = 12.0

_V31_CALC_CACHE = {}
_V31_CACHE_ORDER = []

_V31_LAST_CALLBACK = {}
_V31_LAST_SCREEN = {}
_V31_EXPORT_FILES = {}

# ============================================================
# TIME
# ============================================================

def v31_now():
    return (
        datetime.utcnow()
        .replace(microsecond=0)
        .isoformat()
        + "Z"
    )


# ============================================================
# SAFE NAME
# ============================================================

def v31_safe_name(
    value,
    fallback="project"
):
    text = str(value or "").strip()

    text = re.sub(
        r"[^\w\-\.\u0600-\u06ff ]+",
        "_",
        text,
        flags=re.UNICODE
    )

    text = re.sub(
        r"\s+",
        "_",
        text
    )

    text = text.strip("._")

    return text[:80] or fallback


# ============================================================
# USER ID
# ============================================================

def v31_user_id(context):
    try:
        return str(context._user_id)
    except Exception:
        pass

    try:
        return str(
            context.user_data.get(
                "_user_id",
                "anonymous"
            )
        )
    except Exception:
        return "anonymous"


# ============================================================
# PROJECT ID
# ============================================================

def v31_project_id(context):

    pid = context.user_data.get(
        "project_id"
    )

    if not pid:

        seed = (
            f"{v31_user_id(context)}:"
            f"{time.time_ns()}"
        )

        pid = hashlib.sha1(
            seed.encode("utf-8")
        ).hexdigest()[:12]

        context.user_data[
            "project_id"
        ] = pid

    return str(pid)


# ============================================================
# PROJECT NAME
# ============================================================

def v31_project_name(context):

    return str(
        context.user_data.get(
            "project_name"
        )
        or
        f"Project_{v31_project_id(context)}"
    )


# ============================================================
# PROJECT PATH
# ============================================================

def v31_project_path(
    context,
    project_id=None
):

    uid = v31_safe_name(
        v31_user_id(context),
        "user"
    )

    pid = v31_safe_name(
        project_id or v31_project_id(context),
        "project"
    )

    folder = V31_ROOT / uid

    folder.mkdir(
        parents=True,
        exist_ok=True
    )

    return (
        folder /
        f"{pid}.json"
    )


# ============================================================
# DEEP COPY
# ============================================================

def v31_deepcopy(value):

    try:
        return copy.deepcopy(value)

    except Exception:

        return json.loads(
            json.dumps(
                value,
                ensure_ascii=False,
                default=str
            )
        )


# ============================================================
# JSON SAFE
# ============================================================

def v31_jsonable(value):

    if isinstance(value, dict):

        return {
            str(k): v31_jsonable(v)
            for k, v in value.items()
        }

    if isinstance(
        value,
        (list, tuple)
    ):

        return [
            v31_jsonable(x)
            for x in value
        ]

    if isinstance(value, set):

        return [
            v31_jsonable(x)
            for x in sorted(
                value,
                key=str
            )
        ]

    if (
        isinstance(value, float)
        and not math.isfinite(value)
    ):

        return None

    if isinstance(
        value,
        (
            str,
            int,
            float,
            bool
        )
    ) or value is None:

        return value

    return str(value)


# ============================================================
# PROJECT PAYLOAD
# ============================================================

def v31_project_payload(context):

    data = v31_deepcopy(
        dict(
            context.user_data
        )
    )

    transient = {
        "history",
        "step_index",
        "roof_step_index",
        "awaiting_fy",
        "natural_mode",
        "multi_mode",
        "equiv_step",
        "equiv",
        "manual_input",
        "_v31_busy",
        "_v31_last_callback",
    }

    for key in transient:
        data.pop(
            key,
            None
        )

    data["project_id"] = (
        v31_project_id(context)
    )

    data["project_name"] = (
        v31_project_name(context)
    )

    data["saved_at"] = v31_now()

    data["schema_version"] = (
        V31_VERSION
    )

    return v31_jsonable(
        data
    )


# ============================================================
# ATOMIC WRITE
# ============================================================

def v31_atomic_write(
    path,
    payload
):

    path = Path(path)

    path.parent.mkdir(
        parents=True,
        exist_ok=True
    )

    fd, tmp = tempfile.mkstemp(
        prefix=".project_",
        suffix=".tmp",
        dir=str(path.parent)
    )

    try:

        with os.fdopen(
            fd,
            "w",
            encoding="utf-8"
        ) as fh:

            json.dump(
                payload,
                fh,
                ensure_ascii=False,
                indent=2,
                default=str
            )

            fh.flush()

            os.fsync(
                fh.fileno()
            )

        os.replace(
            tmp,
            path
        )

    finally:

        if os.path.exists(tmp):

            try:
                os.unlink(tmp)
            except OSError:
                pass


# ============================================================
# SAVE PROJECT
# ============================================================

def v31_save_project(
    context,
    reason="manual"
):

    payload = v31_project_payload(
        context
    )

    payload["save_reason"] = reason

    path = v31_project_path(
        context
    )

    v31_atomic_write(
        path,
        payload
    )

    context.user_data[
        "last_saved_at"
    ] = payload["saved_at"]

    context.user_data[
        "project_dirty"
    ] = False

    return path


# ============================================================
# MARK DIRTY
# ============================================================

def v31_mark_dirty(context):

    context.user_data[
        "project_dirty"
    ] = True

    context.user_data[
        "updated_at"
    ] = v31_now()


# ============================================================
# LOAD PROJECT
# ============================================================

def v31_load_project(
    context,
    payload
):

    if not isinstance(
        payload,
        dict
    ):
        raise ValueError(
            "Saved project is not a valid object."
        )

    allowed = v31_jsonable(
        payload
    )

    context.user_data.clear()

    context.user_data.update(
        allowed
    )

    context.user_data[
        "project_dirty"
    ] = False

    context.user_data[
        "loaded_at"
    ] = v31_now()

    context.user_data.pop(
        "save_reason",
        None
    )

    return context.user_data


# ============================================================
# LIST PROJECTS
# ============================================================

def v31_list_projects(context):

    uid = v31_safe_name(
        v31_user_id(context),
        "user"
    )

    folder = (
        V31_ROOT /
        uid
    )

    folder.mkdir(
        parents=True,
        exist_ok=True
    )

    rows = []

    for path in folder.glob(
        "*.json"
    ):

        try:

            stat = path.stat()

            if (
                stat.st_size
                >
                V31_FILE_LIMIT_MB * 1024 * 1024
            ):
                continue

            with path.open(
                "r",
                encoding="utf-8"
            ) as fh:

                data = json.load(
                    fh
                )

            rows.append(
                {
                    "project_id":
                        data.get(
                            "project_id",
                            path.stem
                        ),

                    "project_name":
                        data.get(
                            "project_name",
                            path.stem
                        ),

                    "saved_at":
                        data.get(
                            "saved_at",
                            ""
                        ),

                    "members":
                        len(
                            data.get(
                                "project_results",
                                []
                            )
                            or []
                        ),

                    "path":
                        str(path),

                    "mtime":
                        stat.st_mtime,
                }
            )

        except Exception as exc:

            logger.warning(
                "Skipping invalid project file %s: %s",
                path,
                exc
            )

    rows.sort(
        key=lambda x:
            x.get(
                "mtime",
                0
            ),
        reverse=True
    )

    return rows[
        :V31_PROJECT_LIMIT
    ]


# ============================================================
# COPY PROJECT
# ============================================================

def v31_copy_project_context(
    context,
    new_name=None
):

    source = v31_project_payload(
        context
    )

    copied = v31_deepcopy(
        source
    )

    seed = (
        f"copy:"
        f"{v31_project_id(context)}:"
        f"{time.time_ns()}"
    )

    new_id = hashlib.sha1(
        seed.encode(
            "utf-8"
        )
    ).hexdigest()[:12]

    copied[
        "project_id"
    ] = new_id

    copied[
        "project_name"
    ] = (
        new_name
        or
        f"{v31_project_name(context)}_copy"
    )

    copied[
        "copied_from"
    ] = v31_project_id(
        context
    )

    copied[
        "saved_at"
    ] = v31_now()

    new_path = v31_project_path(
        context,
        new_id
    )

    v31_atomic_write(
        new_path,
        copied
    )

    return (
        new_path,
        copied
    )


# ============================================================
# MEMBER ACCESS
# ============================================================

def v31_member_results(context):

    return context.user_data.setdefault(
        "project_results",
        []
    )


# ============================================================
# COPY MEMBER
# ============================================================

def v31_copy_member(
    context,
    index
):

    rows = v31_member_results(
        context
    )

    index = int(index)

    if (
        index < 0
        or
        index >= len(rows)
    ):
        raise IndexError(
            "Member index is out of range."
        )

    if (
        len(rows)
        >=
        V31_MEMBER_LIMIT
    ):
        raise ValueError(
            "Project member limit reached."
        )

    item = v31_deepcopy(
        rows[index]
    )

    item[
        "copied_from_index"
    ] = index + 1

    item[
        "copy_created_at"
    ] = v31_now()

    base = str(
        item.get(
            "member_name"
        )
        or
        item.get(
            "member_type"
        )
        or
        item.get(
            "kind"
        )
        or
        "Member"
    )

    item[
        "member_name"
    ] = (
        f"{base} - Copy"
    )

    rows.append(
        item
    )

    v31_mark_dirty(
        context
    )

    return (
        len(rows) - 1,
        item
    )


# ============================================================
# DELETE MEMBER
# ============================================================

def v31_delete_member(
    context,
    index
):

    rows = v31_member_results(
        context
    )

    index = int(index)

    if (
        index < 0
        or
        index >= len(rows)
    ):
        raise IndexError(
            "Member index is out of range."
        )

    deleted = rows.pop(
        index
    )

    v31_mark_dirty(
        context
    )

    return deleted


# ============================================================
# QUICK INPUT VALUES
# ============================================================

V31_DIAMETERS = (
    8,
    10,
    12,
    14,
    16,
    18,
    20,
    22,
    25,
    28,
    32,
)

V31_SPACINGS = (
    10,
    12,
    15,
    20,
    25,
    30,
    40,
    50,
    75,
    100,
    125,
    150,
    200,
)

V31_COUNTS = (
    1,
    2,
    3,
    4,
    5,
    6,
    8,
    10,
    12,
    16,
    20,
    24,
)

V31_LENGTHS = (
    0.10,
    0.15,
    0.20,
    0.25,
    0.30,
    0.40,
    0.50,
    0.60,
    0.80,
    1.00,
    1.50,
    2.00,
    3.00,
    4.00,
    5.00,
    6.00,
)


# ============================================================
# QUICK INPUT TYPE
# ============================================================

def v31_preset_kind(
    field_type,
    key
):

    k = str(
        key
    ).lower()

    t = str(
        field_type
    ).lower()

    if (
        "dia" in k
        or
        "diameter" in t
    ):
        return "diameter"

    if (
        "spacing" in k
        or
        "spacing" in t
    ):
        return "spacing"

    if (
        t in (
            "int",
            "int0"
        )
        or
        k in {
            "count",
            "lc",
            "tlc",
            "floors",
            "steps",
        }
    ):
        return "count"

    if t == "float":
        return "length"

    return None


# ============================================================
# QUICK BUTTON ROWS
# ============================================================

def v31_preset_rows(
    field_type,
    key,
    prefix="v31preset"
):

    kind = v31_preset_kind(
        field_type,
        key
    )

    values = {
        "diameter":
            V31_DIAMETERS,

        "spacing":
            V31_SPACINGS,

        "count":
            V31_COUNTS,

        "length":
            V31_LENGTHS,
    }.get(
        kind,
        ()
    )

    rows = []

    for start in range(
        0,
        len(values),
        4
    ):

        row = []

        for value in values[
            start:start + 4
        ]:

            if kind == "diameter":

                label = (
                    f"Φ{value}"
                )

            elif kind == "length":

                label = (
                    f"{value:g} m"
                )

            elif kind == "spacing":

                label = (
                    f"{value} mm"
                )

            else:

                label = str(
                    value
                )

            row.append(
                (
                    label,
                    f"{prefix}_{value}"
                )
            )

        rows.append(
            row
        )

    return rows


# ============================================================
# CALCULATION CACHE KEY
# ============================================================

def v31_calc_key(
    kind,
    values,
    project_settings
):

    blob = json.dumps(
        {
            "kind": kind,
            "values":
                v31_jsonable(values),
            "settings":
                v31_jsonable(
                    project_settings
                ),
        },
        ensure_ascii=False,
        sort_keys=True,
        separators=(
            ",",
            ":"
        )
    )

    return hashlib.sha256(
        blob.encode(
            "utf-8"
        )
    ).hexdigest()


# ============================================================
# CACHE GET
# ============================================================

def v31_cache_get(key):

    item = (
        _V31_CALC_CACHE.get(
            key
        )
    )

    if item is None:
        return None

    _V31_CACHE_ORDER.append(
        key
    )

    return v31_deepcopy(
        item
    )


# ============================================================
# CACHE PUT
# ============================================================

def v31_cache_put(
    key,
    value
):

    _V31_CALC_CACHE[
        key
    ] = v31_deepcopy(
        value
    )

    _V31_CACHE_ORDER.append(
        key
    )

    while (
        len(_V31_CALC_CACHE)
        >
        V31_CALC_CACHE_LIMIT
    ):

        while _V31_CACHE_ORDER:

            old = (
                _V31_CACHE_ORDER.pop(
                    0
                )
            )

            if (
                old
                in
                _V31_CALC_CACHE
                and
                old != key
            ):

                del _V31_CALC_CACHE[
                    old
                ]

                break

        else:

            break


# ============================================================
# METRICS
# ============================================================

def v31_metrics(context):

    try:

        return project_metrics(
            context
        )

    except Exception:

        rows = v31_member_results(
            context
        )

        concrete = sum(
            _safe_float(
                x.get(
                    "concrete_m3",
                    x.get(
                        "concrete",
                        0
                    )
                )
            )
            for x in rows
        )

        steel = sum(
            _safe_float(
                x.get(
                    "rebar_kg",
                    x.get(
                        "steel_kg",
                        0
                    )
                )
            )
            for x in rows
        )

        return {
            "members":
                len(rows),

            "concrete_m3":
                concrete,

            "steel_kg":
                steel,
        }


# ============================================================
# ALL REBAR
# ============================================================

def v31_all_rebar(context):

    out = []

    for i, member in enumerate(
        v31_member_results(context),
        1
    ):

        details = (
            member.get(
                "rebar_details"
            )
            or
            member.get(
                "details"
            )
            or
            []
        )

        if isinstance(
            details,
            dict
        ):

            details = [
                details
            ]

        for j, d in enumerate(
            details,
            1
        ):

            if not isinstance(
                d,
                dict
            ):
                continue

            row = dict(
                d
            )

            row[
                "member_index"
            ] = i

            row[
                "member_name"
            ] = (
                member.get(
                    "member_name"
                )
                or
                member.get(
                    "member_type"
                )
                or
                member.get(
                    "kind"
                )
                or
                f"Member {i}"
            )

            row[
                "bar_mark"
            ] = (
                row.get(
                    "bar_mark"
                )
                or
                row.get(
                    "mark"
                )
                or
                f"M{i}-B{j}"
            )

            row[
                "diameter_mm"
            ] = _safe_int(
                row.get(
                    "diameter_mm",
                    row.get(
                        "dia_mm",
                        row.get(
                            "diameter",
                            0
                        )
                    )
                )
            )

            row[
                "length_m"
            ] = _safe_float(
                row.get(
                    "length_m",
                    row.get(
                        "length",
                        0
                    )
                )
            )

            row[
                "count"
            ] = _safe_int(
                row.get(
                    "count",
                    row.get(
                        "quantity",
                        0
                    )
                )
            )

            row[
                "weight_kg"
            ] = _safe_float(
                row.get(
                    "weight_kg",
                    row.get(
                        "weight",
                        0
                    )
                )
            )

            out.append(
                row
            )

    return out


# ============================================================
# CUT PLAN
# ============================================================

def v31_cut_plan(context):

    bars = v31_all_rebar(
        context
    )

    pieces = []

    for row in bars:

        length = _safe_float(
            row.get(
                "length_m"
            )
        )

        count = max(
            0,
            _safe_int(
                row.get(
                    "count"
                )
            )
        )

        if (
            length <= 0
            or
            count <= 0
        ):
            continue

        if (
            length
            >
            V31_STOCK_LENGTH
            + 1e-9
        ):

            pieces.append(
                {
                    **row,

                    "cut_status":
                        "OVER_12M",

                    "stock_bars":
                        count,

                    "waste_m":
                        0.0,
                }
            )

            continue

        stock = (
            math.ceil(
                length /
                V31_STOCK_LENGTH
            )
            *
            count
        )

        waste = max(
            stock *
            V31_STOCK_LENGTH
            -
            length * count,
            0.0
        )

        pieces.append(
            {
                **row,

                "cut_status":
                    "OK",

                "stock_bars":
                    stock,

                "waste_m":
                    waste,
            }
        )

    return pieces


# ============================================================
# REPORT ROWS
# ============================================================

def v31_report_rows(context):

    rows = []

    for i, r in enumerate(
        v31_member_results(context),
        1
    ):

        rows.append(
            {
                "#": i,

                "member":
                    r.get(
                        "member_name"
                    )
                    or
                    r.get(
                        "member_type"
                    )
                    or
                    r.get(
                        "kind"
                    )
                    or
                    f"Member {i}",

                "type":
                    r.get(
                        "member_type"
                    )
                    or
                    r.get(
                        "kind"
                    )
                    or
                    "",

                "concrete_m3":
                    _safe_float(
                        r.get(
                            "concrete_m3",
                            r.get(
                                "concrete",
                                0
                            )
                        )
                    ),

                "rebar_kg":
                    _safe_float(
                        r.get(
                            "rebar_kg",
                            r.get(
                                "steel_kg",
                                0
                            )
                        )
                    ),

                "area_m2":
                    _safe_float(
                        r.get(
                            "area_m2",
                            r.get(
                                "net_area_m2",
                                0
                            )
                        )
                    ),
            }
        )

    return rows


# ============================================================
# EXCEL EXPORT
# ============================================================

def v31_make_excel(context):

    try:

        from openpyxl import Workbook

        from openpyxl.styles import (
            Font,
            Alignment,
            Border,
            Side
        )

        from openpyxl.utils import (
            get_column_letter
        )

    except ImportError as exc:

        raise RuntimeError(
            "openpyxl is required for Excel export."
        ) from exc

    wb = Workbook()

    ws = wb.active

    ws.title = (
        "Project Summary"
    )

    ws.append(
        [
            "STRUCTURAL QUANTITY ENGINE"
        ]
    )

    ws.append(
        [
            "Version",
            V31_VERSION
        ]
    )

    ws.append(
        [
            "Project",
            v31_project_name(
                context
            )
        ]
    )

    ws.append(
        [
            "Project ID",
            v31_project_id(
                context
            )
        ]
    )

    ws.append(
        [
            "Standard",
            context.user_data.get(
                "standard",
                ""
            )
        ]
    )

    ws.append(
        [
            "Concrete",
            f"C{context.user_data.get('fc', '')}"
        ]
    )

    ws.append(
        [
            "Rebar",
            context.user_data.get(
                "rebar_grade",
                ""
            )
        ]
    )

    ws.append(
        [
            "fy (MPa)",
            context.user_data.get(
                "fy",
                ""
            )
        ]
    )

    ws.append(
        [
            "Generated",
            v31_now()
        ]
    )

    ws.append([])

    m = v31_metrics(
        context
    )

    ws.append(
        [
            "Metric",
            "Value",
            "Unit"
        ]
    )

    ws.append(
        [
            "Members",
            m.get(
                "members",
                0
            ),
            "ea"
        ]
    )

    ws.append(
        [
            "Concrete",
            m.get(
                "concrete_m3",
                0
            ),
            "m³"
        ]
    )

    ws.append(
        [
            "Rebar",
            m.get(
                "steel_kg",
                0
            ),
            "kg"
        ]
    )

    ws.append(
        [
            "Area",
            m.get(
                "area_m2",
                0
            ),
            "m²"
        ]
    )

    thin = Side(
        style="thin"
    )

    for cell in ws[1]:

        cell.font = Font(
            bold=True,
            size=14
        )

    for row in ws.iter_rows():

        for cell in row:

            cell.alignment = Alignment(
                vertical="top"
            )

    for col in range(
        1,
        5
    ):

        ws.column_dimensions[
            get_column_letter(col)
        ].width = 22

    members = wb.create_sheet(
        "Members"
    )

    headers = [
        "#",
        "Member",
        "Type",
        "Concrete m3",
        "Rebar kg",
        "Area m2"
    ]

    members.append(
        headers
    )

    for row in v31_report_rows(
        context
    ):

        members.append(
            [
                row[h]
                for h in headers
            ]
        )

    for cell in members[1]:

        cell.font = Font(
            bold=True
        )

        cell.border = Border(
            bottom=thin
        )

    bars = wb.create_sheet(
        "Bar Schedule"
    )

    bar_headers = [
        "Member",
        "Bar Mark",
        "Diameter mm",
        "Count",
        "Length m",
        "Weight kg",
        "Location"
    ]

    bars.append(
        bar_headers
    )

    for row in v31_all_rebar(
        context
    ):

        bars.append(
            [
                row.get(
                    "member_name",
                    ""
                ),

                row.get(
                    "bar_mark",
                    ""
                ),

                row.get(
                    "diameter_mm",
                    0
                ),

                row.get(
                    "count",
                    0
                ),

                row.get(
                    "length_m",
                    0
                ),

                row.get(
                    "weight_kg",
                    0
                ),

                row.get(
                    "location",
                    row.get(
                        "zone",
                        ""
                    )
                ),
            ]
        )

    for cell in bars[1]:

        cell.font = Font(
            bold=True
        )

        cell.border = Border(
            bottom=thin
        )

    cut = wb.create_sheet(
        "Cut List"
    )

    cut_headers = [
        "Member",
        "Bar Mark",
        "Diameter mm",
        "Piece m",
        "Count",
        "12m Stock Bars",
        "Waste m",
        "Status"
    ]

    cut.append(
        cut_headers
    )

    for row in v31_cut_plan(
        context
    ):

        cut.append(
            [
                row.get(
                    "member_name",
                    ""
                ),

                row.get(
                    "bar_mark",
                    ""
                ),

                row.get(
                    "diameter_mm",
                    0
                ),

                row.get(
                    "length_m",
                    0
                ),

                row.get(
                    "count",
                    0
                ),

                row.get(
                    "stock_bars",
                    0
                ),

                row.get(
                    "waste_m",
                    0
                ),

                row.get(
                    "cut_status",
                    ""
                ),
            ]
        )

    for cell in cut[1]:

        cell.font = Font(
            bold=True
        )

        cell.border = Border(
            bottom=thin
        )

    qa = wb.create_sheet(
        "QA"
    )

    qa.append(
        [
            "Check",
            "Status",
            "Details"
        ]
    )

    try:

        results = run_project_qa(
            context
        )

        for item in results:

            qa.append(
                [
                    item.get(
                        "id",
                        ""
                    ),

                    (
                        "PASS"
                        if item.get(
                            "ok"
                        )
                        else
                        "FAIL"
                    ),

                    "; ".join(
                        item.get(
                            "issues",
                            []
                        )
                    ),
                ]
            )

    except Exception as exc:

        qa.append(
            [
                "QA",
                "ERROR",
                str(exc)
            ]
        )

    for sheet in wb.worksheets:

        for column_cells in sheet.columns:

            max_len = 0

            for cell in column_cells:

                try:

                    max_len = max(
                        max_len,
                        len(
                            str(
                                cell.value
                                or
                                ""
                            )
                        )
                    )

                except Exception:
                    pass

            sheet.column_dimensions[
                get_column_letter(
                    column_cells[
                        0
                    ].column
                )
            ].width = min(
                max(
                    max_len + 2,
                    12
                ),
                42
            )

        sheet.freeze_panes = "A2"

    path = (
        V31_ROOT /
        v31_safe_name(
            v31_user_id(context),
            "user"
        )
    )

    path.mkdir(
        parents=True,
        exist_ok=True
    )

    filename = (
        f"{v31_safe_name(v31_project_name(context))}_"
        f"{v31_project_id(context)}.xlsx"
    )

    target = (
        path /
        filename
    )

    wb.save(
        target
    )

    return target


# ============================================================
# PDF EXPORT
# ============================================================

def v31_make_pdf(context):

    try:

        from reportlab.lib import colors

        from reportlab.lib.pagesizes import A4

        from reportlab.lib.styles import (
            getSampleStyleSheet,
            ParagraphStyle
        )

        from reportlab.lib.enums import (
            TA_CENTER
        )

        from reportlab.platypus import (
            SimpleDocTemplate,
            Paragraph,
            Spacer,
            Table,
            TableStyle,
            PageBreak
        )

    except ImportError as exc:

        raise RuntimeError(
            "reportlab is required for PDF export."
        ) from exc

    path = (
        V31_ROOT /
        v31_safe_name(
            v31_user_id(context),
            "user"
        )
    )

    path.mkdir(
        parents=True,
        exist_ok=True
    )

    target = (
        path /
        (
            f"{v31_safe_name(v31_project_name(context))}_"
            f"{v31_project_id(context)}.pdf"
        )
    )

    doc = SimpleDocTemplate(
        str(target),
        pagesize=A4,
        rightMargin=30,
        leftMargin=30,
        topMargin=30,
        bottomMargin=30
    )

    styles = getSampleStyleSheet()

    styles.add(
        ParagraphStyle(
            name="V31Center",
            parent=styles["Title"],
            alignment=TA_CENTER,
            fontSize=15
        )
    )

    story = []

    story.append(
        Paragraph(
            "STRUCTURAL QUANTITY ENGINE",
            styles["V31Center"]
        )
    )

    story.append(
        Spacer(
            1,
            10
        )
    )

    info = [
        [
            "Project",
            v31_project_name(
                context
            )
        ],

        [
            "Project ID",
            v31_project_id(
                context
            )
        ],

        [
            "Standard",
            context.user_data.get(
                "standard",
                ""
            )
        ],

        [
            "Concrete",
            f"C{context.user_data.get('fc', '')}"
        ],

        [
            "Rebar",
            context.user_data.get(
                "rebar_grade",
                ""
            )
        ],

        [
            "fy",
            str(
                context.user_data.get(
                    "fy",
                    ""
                )
            )
            +
            " MPa"
        ],

        [
            "Generated",
            v31_now()
        ],
    ]

    table = Table(
        info,
        colWidths=[
            100,
            390
        ]
    )

    table.setStyle(
        TableStyle(
            [
                (
                    "GRID",
                    (0, 0),
                    (-1, -1),
                    0.4,
                    colors.grey
                ),

                (
                    "FONTNAME",
                    (0, 0),
                    (0, -1),
                    "Helvetica-Bold"
                ),

                (
                    "VALIGN",
                    (0, 0),
                    (-1, -1),
                    "TOP"
                ),
            ]
        )
    )

    story.append(
        table
    )

    story.append(
        Spacer(
            1,
            14
        )
    )

    m = v31_metrics(
        context
    )

    summary = [
        [
            "Metric",
            "Value",
            "Unit"
        ],

        [
            "Members",
            str(
                m.get(
                    "members",
                    0
                )
            ),
            "ea"
        ],

        [
            "Concrete",
            f"{m.get('concrete_m3', 0):.3f}",
            "m3"
        ],

        [
            "Rebar",
            f"{m.get('steel_kg', 0):.1f}",
            "kg"
        ],

        [
            "Area",
            f"{m.get('area_m2', 0):.3f}",
            "m2"
        ],
    ]

    t2 = Table(
        summary,
        colWidths=[
            180,
            180,
            130
        ]
    )

    t2.setStyle(
        TableStyle(
            [
                (
                    "GRID",
                    (0, 0),
                    (-1, -1),
                    0.4,
                    colors.grey
                ),

                (
                    "BACKGROUND",
                    (0, 0),
                    (-1, 0),
                    colors.lightgrey
                ),

                (
                    "FONTNAME",
                    (0, 0),
                    (-1, 0),
                    "Helvetica-Bold"
                ),
            ]
        )
    )

    story.append(
        t2
    )

    story.append(
        Spacer(
            1,
            16
        )
    )

    member_data = [
        [
            "#",
            "Member",
            "Type",
            "Concrete",
            "Rebar kg"
        ]
    ]

    for row in v31_report_rows(
        context
    ):

        member_data.append(
            [
                str(
                    row["#"]
                ),

                str(
                    row["member"]
                )[:35],

                str(
                    row["type"]
                )[:20],

                f"{row['concrete_m3']:.3f}",

                f"{row['rebar_kg']:.1f}",
            ]
        )

    if len(
        member_data
    ) == 1:

        member_data.append(
            [
                "-",
                "No calculated members",
                "-",
                "0",
                "0"
            ]
        )

    mt = Table(
        member_data,
        repeatRows=1,
        colWidths=[
            25,
            190,
            90,
            90,
            90
        ]
    )

    mt.setStyle(
        TableStyle(
            [
                (
                    "GRID",
                    (0, 0),
                    (-1, -1),
                    0.35,
                    colors.grey
                ),

                (
                    "BACKGROUND",
                    (0, 0),
                    (-1, 0),
                    colors.lightgrey
                ),

                (
                    "FONTNAME",
                    (0, 0),
                    (-1, 0),
                    "Helvetica-Bold"
                ),

                (
                    "FONTSIZE",
                    (0, 0),
                    (-1, -1),
                    7.5
                ),

                (
                    "VALIGN",
                    (0, 0),
                    (-1, -1),
                    "TOP"
                ),
            ]
        )
    )

    story.append(
        Paragraph(
            "Members",
            styles["Heading2"]
        )
    )

    story.append(
        mt
    )

    story.append(
        PageBreak()
    )

    bars_data = [
        [
            "Member",
            "Bar Mark",
            "Dia",
            "Count",
            "Length",
            "Weight"
        ]
    ]

    for row in v31_all_rebar(
        context
    ):

        bars_data.append(
            [
                str(
                    row.get(
                        "member_name",
                        ""
                    )
                )[:24],

                str(
                    row.get(
                        "bar_mark",
                        ""
                    )
                ),

                str(
                    row.get(
                        "diameter_mm",
                        0
                    )
                ),

                str(
                    row.get(
                        "count",
                        0
                    )
                ),

                f"{_safe_float(row.get('length_m')):.3f}",

                f"{_safe_float(row.get('weight_kg')):.2f}",
            ]
        )

    if len(
        bars_data
    ) == 1:

        bars_data.append(
            [
                "-",
                "-",
                "-",
                "0",
                "0",
                "0"
            ]
        )

    bt = Table(
        bars_data,
        repeatRows=1,
        colWidths=[
            150,
            75,
            45,
            50,
            75,
            75
        ]
    )

    bt.setStyle(
        TableStyle(
            [
                (
                    "GRID",
                    (0, 0),
                    (-1, -1),
                    0.3,
                    colors.grey
                ),

                (
                    "BACKGROUND",
                    (0, 0),
                    (-1, 0),
                    colors.lightgrey
                ),

                (
                    "FONTNAME",
                    (0, 0),
                    (-1, 0),
                    "Helvetica-Bold"
                ),

                (
                    "FONTSIZE",
                    (0, 0),
                    (-1, -1),
                    7.2
                ),
            ]
        )
    )

    story.append(
        Paragraph(
            "Bar Schedule",
            styles["Heading2"]
        )
    )

    story.append(
        bt
    )

    story.append(
        Spacer(
            1,
            12
        )
    )

    cut_data = [
        [
            "Member",
            "Mark",
            "Dia",
            "Piece",
            "Count",
            "12m",
            "Waste"
        ]
    ]

    for row in v31_cut_plan(
        context
    ):

        cut_data.append(
            [
                str(
                    row.get(
                        "member_name",
                        ""
                    )
                )[:20],

                str(
                    row.get(
                        "bar_mark",
                        ""
                    )
                ),

                str(
                    row.get(
                        "diameter_mm",
                        0
                    )
                ),

                f"{_safe_float(row.get('length_m')):.3f}",

                str(
                    row.get(
                        "count",
                        0
                    )
                ),

                str(
                    row.get(
                        "stock_bars",
                        0
                    )
                ),

                f"{_safe_float(row.get('waste_m')):.3f}",
            ]
        )

    if len(
        cut_data
    ) == 1:

        cut_data.append(
            [
                "-",
                "-",
                "-",
                "0",
                "0",
                "0",
                "0"
            ]
        )

    ct = Table(
        cut_data,
        repeatRows=1,
        colWidths=[
            130,
            55,
            40,
            60,
            45,
            45,
            60
        ]
    )

    ct.setStyle(
        TableStyle(
            [
                (
                    "GRID",
                    (0, 0),
                    (-1, -1),
                    0.3,
                    colors.grey
                ),

                (
                    "BACKGROUND",
                    (0, 0),
                    (-1, 0),
                    colors.lightgrey
                ),

                (
                    "FONTNAME",
                    (0, 0),
                    (-1, 0),
                    "Helvetica-Bold"
                ),

                (
                    "FONTSIZE",
                    (0, 0),
                    (-1, -1),
                    7.0
                ),
            ]
        )
    )

    story.append(
        Paragraph(
            "Cut List / 12m Stock",
            styles["Heading2"]
        )
    )

    story.append(
        ct
    )

    story.append(
        Spacer(
            1,
            16
        )
    )

    story.append(
        Paragraph(
            "Engineering note: quantity/detailing output must be checked against approved structural drawings and the governing code.",
            styles["BodyText"]
        )
    )

    doc.build(
        story
    )

    return target


# ============================================================
# JSON EXPORT
# ============================================================

def v31_make_json(context):

    path = (
        V31_ROOT /
        v31_safe_name(
            v31_user_id(context),
            "user"
        )
    )

    path.mkdir(
        parents=True,
        exist_ok=True
    )

    target = (
        path /
        (
            f"{v31_safe_name(v31_project_name(context))}_"
            f"{v31_project_id(context)}.json"
        )
    )

    with target.open(
        "w",
        encoding="utf-8"
    ) as fh:

        json.dump(
            v31_project_payload(
                context
            ),
            fh,
            ensure_ascii=False,
            indent=2
        )

    return target


# ============================================================
# CSV EXPORT
# ============================================================

def v31_make_csv(context):

    path = (
        V31_ROOT /
        v31_safe_name(
            v31_user_id(context),
            "user"
        )
    )

    path.mkdir(
        parents=True,
        exist_ok=True
    )

    target = (
        path /
        (
            f"{v31_safe_name(v31_project_name(context))}_"
            f"{v31_project_id(context)}.csv"
        )
    )

    rows = v31_report_rows(
        context
    )

    with target.open(
        "w",
        encoding="utf-8-sig",
        newline=""
    ) as fh:

        writer = csv.DictWriter(
            fh,
            fieldnames=[
                "#",
                "member",
                "type",
                "concrete_m3",
                "rebar_kg",
                "area_m2"
            ]
        )

        writer.writeheader()

        writer.writerows(
            rows
        )

    return target


# ============================================================
# DASHBOARD
# ============================================================

def v31_dashboard_text(
    context,
    lang="fa"
):

    m = v31_metrics(
        context
    )

    name = escape(
        v31_project_name(
            context
        )
    )

    pid = escape(
        v31_project_id(
            context
        )
    )

    dirty = (
        "⚠️ ذخیره نشده"
        if context.user_data.get(
            "project_dirty"
        )
        else
        "💾 ذخیره شده"
    )

    return (
        "📊 <b>داشبورد پروژه</b>\n\n"

        f"📁 نام: <b>{name}</b>\n"

        f"🆔 ID: <code>{pid}</code>\n"

        f"{dirty}\n\n"

        f"👷 اعضا: "
        f"{m.get('members', 0)}\n"

        f"🧱 بتن: "
        f"{m.get('concrete_m3', 0):.3f} m³\n"

        f"🔩 میلگرد: "
        f"{m.get('steel_kg', 0):.1f} kg\n"

        f"📐 مساحت: "
        f"{m.get('area_m2', 0):.3f} m²\n\n"

        f"📚 استاندارد: "
        f"{escape(str(context.user_data.get('standard', '-')))}\n"

        f"🧱 بتن: "
        f"C{context.user_data.get('fc', '-')}\n"

        f"🔩 میلگرد: "
        f"{escape(str(context.user_data.get('rebar_grade', '-')))}"
    )


# ============================================================
# PROJECT LIST TEXT
# ============================================================

def v31_project_list_text(
    rows
):

    if not rows:

        return (
            "💾 <b>پروژه‌های ذخیره‌شده</b>\n\n"
            "هنوز پروژه‌ای ذخیره نشده است."
        )

    lines = [
        "💾 <b>پروژه‌های ذخیره‌شده</b>",
        ""
    ]

    for i, row in enumerate(
        rows[:20],
        1
    ):

        lines.append(
            (
                f"{i}. "
                f"<b>{escape(str(row['project_name']))}</b>"
                f" | {row['members']} عضو"
                f" | {escape(str(row['saved_at']))}"
            )
        )

    return "\n".join(
        lines
    )


# ============================================================
# PROJECT KEYBOARD
# ============================================================

def v31_project_keyboard(
    lang="fa"
):

    return kb(
        [
            [
                (
                    "💾 ذخیره پروژه",
                    "v31_save"
                ),

                (
                    "📂 پروژه‌ها",
                    "v31_projects"
                ),
            ],

            [
                (
                    "📋 کپی پروژه",
                    "v31_copy_project"
                ),

                (
                    "📊 داشبورد",
                    "v31_dashboard"
                ),
            ],

            [
                (
                    "📥 Excel",
                    "v31_export_xlsx"
                ),

                (
                    "📄 PDF",
                    "v31_export_pdf"
                ),
            ],

            [
                (
                    "🗂 JSON",
                    "v31_export_json"
                ),

                (
                    "📑 CSV",
                    "v31_export_csv"
                ),
            ],

            [
                (
                    "🏠 صفحه اصلی",
                    "home"
                )
            ],
        ]
    )


# ============================================================
# RESULT KEYBOARD
# ============================================================

def v31_result_keyboard(
    lang="fa"
):

    return kb(
        [
            [
                (
                    "🔩 جزئیات میلگرد",
                    "show_rebar"
                ),

                (
                    "✂️ Cut List",
                    "show_cut"
                ),
            ],

            [
                (
                    "📋 کپی عضو",
                    "v31_copy_last_member"
                ),

                (
                    "✏️ ویرایش",
                    "edit_member"
                ),
            ],

            [
                (
                    "💾 ذخیره",
                    "v31_save"
                ),

                (
                    "📥 خروجی",
                    "v31_export_menu"
                ),
            ],

            [
                (
                    "➕ عضو جدید",
                    "new_member"
                ),

                (
                    "📊 خلاصه",
                    "summary"
                ),
            ],

            [
                (
                    "🏠 خانه",
                    "home"
                )
            ],
        ]
    )


# ============================================================
# EXPORT KEYBOARD
# ============================================================

def v31_export_keyboard():

    return kb(
        [
            [
                (
                    "📊 Excel",
                    "v31_export_xlsx"
                ),

                (
                    "📄 PDF",
                    "v31_export_pdf"
                ),
            ],

            [
                (
                    "🗂 JSON",
                    "v31_export_json"
                ),

                (
                    "📑 CSV",
                    "v31_export_csv"
                ),
            ],

            [
                (
                    "📋 پروژه",
                    "v31_dashboard"
                )
            ],
        ]
    )


# ============================================================
# TELEGRAM FILE SENDER
# ============================================================

async def v31_send_file(
    q,
    path,
    caption
):

    path = Path(
        path
    )

    if not path.exists():

        raise FileNotFoundError(
            str(path)
        )

    if (
        path.stat().st_size
        >
        V31_FILE_LIMIT_MB *
        1024 *
        1024
    ):

        raise ValueError(
            "Generated file is larger than configured limit."
        )

    with path.open(
        "rb"
    ) as fh:

        await q.message.reply_document(
            document=fh,
            filename=path.name,
            caption=caption
        )


# ============================================================
# CALLBACK GUARD
# ============================================================

async def v31_ack_guard(
    q,
    context
):

    uid = v31_user_id(
        context
    )

    key = (
        f"{uid}:"
        f"{q.data or ''}"
    )

    now = time.monotonic()

    old = _V31_LAST_CALLBACK.get(
        key,
        0.0
    )

    if (
        now - old
        <
        V31_CALLBACK_GUARD_SECONDS
    ):

        try:

            await q.answer(
                "⏳ در حال پردازش...",
                show_alert=False
            )

        except Exception:
            pass

        return False

    _V31_LAST_CALLBACK[
        key
    ] = now

    if (
        len(_V31_LAST_CALLBACK)
        >
        2000
    ):

        cutoff = (
            now - 30.0
        )

        for k, t in list(
            _V31_LAST_CALLBACK.items()
        ):

            if t < cutoff:

                _V31_LAST_CALLBACK.pop(
                    k,
                    None
                )

    return True


# ============================================================
# SPECIAL CALLBACKS
# ============================================================

async def v31_handle_special_callback(
    update,
    context,
    data
):

    q = update.callback_query

    lang = context.user_data.get(
        "lang",
        "fa"
    )

    # --------------------------------------------------------
    # SAVE
    # --------------------------------------------------------

    if data == "v31_save":

        path = await asyncio.to_thread(
            v31_save_project,
            context,
            "manual"
        )

        await q.edit_message_text(
            (
                "💾 <b>پروژه ذخیره شد</b>\n\n"
                f"📁 {escape(v31_project_name(context))}\n"
                f"🆔 <code>{escape(v31_project_id(context))}</code>"
            ),
            parse_mode="HTML",
            reply_markup=v31_project_keyboard(
                lang
            )
        )

        return True

    # --------------------------------------------------------
    # PROJECT LIST
    # --------------------------------------------------------

    if data == "v31_projects":

        rows = await asyncio.to_thread(
            v31_list_projects,
            context
        )

        buttons = []

        for row in rows[:20]:

            label = (
                f"📁 "
                f"{str(row['project_name'])[:28]}"
            )

            buttons.append(
                [
                    (
                        label,
                        f"v31_load_{row['project_id']}"
                    )
                ]
            )

        buttons.append(
            [
                (
                    "💾 ذخیره فعلی",
                    "v31_save"
                ),

                (
                    "🏠 خانه",
                    "home"
                ),
            ]
        )

        await q.edit_message_text(
            v31_project_list_text(
                rows
            ),
            parse_mode="HTML",
            reply_markup=kb(
                buttons
            )
        )

        return True

    # --------------------------------------------------------
    # LOAD PROJECT
    # --------------------------------------------------------

    if data.startswith(
        "v31_load_"
    ):

        pid = data[
            len("v31_load_"):
        ]

        rows = await asyncio.to_thread(
            v31_list_projects,
            context
        )

        selected = next(
            (
                x
                for x in rows
                if str(
                    x["project_id"]
                ) == pid
            ),
            None
        )

        if not selected:

            await q.answer(
                "پروژه پیدا نشد",
                show_alert=True
            )

            return True

        path = Path(
            selected["path"]
        )

        if (
            path.stat().st_size
            >
            V31_FILE_LIMIT_MB *
            1024 *
            1024
        ):

            raise ValueError(
                "Saved project exceeds configured file limit."
            )

        payload = await asyncio.to_thread(
            lambda:
                json.loads(
                    path.read_text(
                        encoding="utf-8"
                    )
                )
        )

        v31_load_project(
            context,
            payload
        )

        lang = context.user_data.get(
            "lang",
            lang
        )

        await q.edit_message_text(
            v31_dashboard_text(
                context,
                lang
            ),
            parse_mode="HTML",
            reply_markup=v31_project_keyboard(
                lang
            )
        )

        return True

    # --------------------------------------------------------
    # COPY PROJECT
    # --------------------------------------------------------

    if data == "v31_copy_project":

        path, copied = (
            await asyncio.to_thread(
                v31_copy_project_context,
                context
            )
        )

        await q.edit_message_text(
            (
                "📋 <b>کپی پروژه ساخته شد</b>\n\n"
                f"📁 {escape(str(copied['project_name']))}\n"
                f"🆔 <code>{escape(str(copied['project_id']))}</code>"
            ),
            parse_mode="HTML",
            reply_markup=v31_project_keyboard(
                lang
            )
        )

        return True

    # --------------------------------------------------------
    # DASHBOARD
    # --------------------------------------------------------

    if data == "v31_dashboard":

        await q.edit_message_text(
            v31_dashboard_text(
                context,
                lang
            ),
            parse_mode="HTML",
            reply_markup=v31_project_keyboard(
                lang
            )
        )

        return True

    # --------------------------------------------------------
    # COPY LAST MEMBER
    # --------------------------------------------------------

    if data == "v31_copy_last_member":

        rows = v31_member_results(
            context
        )

        if not rows:

            await q.answer(
                "عضوی برای کپی وجود ندارد.",
                show_alert=True
            )

            return True

        idx, item = await asyncio.to_thread(
            v31_copy_member,
            context,
            len(rows) - 1
        )

        await asyncio.to_thread(
            v31_save_project,
            context,
            "member_copy"
        )

        await q.edit_message_text(
            (
                "📋 <b>عضو کپی شد</b>\n\n"
                f"عضو جدید: #{idx + 1}\n"
                f"{escape(str(item.get('member_name', item.get('kind', 'Member'))))}"
            ),
            parse_mode="HTML",
            reply_markup=v31_result_keyboard(
                lang
            )
        )

        return True

    # --------------------------------------------------------
    # EXPORT MENU
    # --------------------------------------------------------

    if data == "v31_export_menu":

        await q.edit_message_text(
            "📥 <b>نوع خروجی را انتخاب کنید</b>",
            parse_mode="HTML",
            reply_markup=v31_export_keyboard()
        )

        return True

    # --------------------------------------------------------
    # EXPORTS
    # --------------------------------------------------------

    if data in {
        "v31_export_xlsx",
        "v31_export_pdf",
        "v31_export_json",
        "v31_export_csv"
    }:

        makers = {

            "v31_export_xlsx":
                (
                    v31_make_excel,
                    "📊 Excel پروژه"
                ),

            "v31_export_pdf":
                (
                    v31_make_pdf,
                    "📄 PDF پروژه"
                ),

            "v31_export_json":
                (
                    v31_make_json,
                    "🗂 JSON پروژه"
                ),

            "v31_export_csv":
                (
                    v31_make_csv,
                    "📑 CSV پروژه"
                ),
        }

        maker, caption = makers[
            data
        ]

        path = await asyncio.to_thread(
            maker,
            context
        )

        await v31_send_file(
            q,
            path,
            caption
        )

        return True

    return False


# ============================================================
# ORIGINAL FUNCTIONS
# ============================================================

_v31_original_do_calculate = (
    do_calculate
)

_v31_original_buttons = (
    buttons
)

_v31_original_main_menu = (
    main_menu
)

_v31_original_result_kb = (
    result_kb
)

# ============================================================
# CALCULATION WRAPPER
# ============================================================

async def v31_do_calculate(
    update,
    context
):

    started = (
        time.perf_counter()
    )

    context.user_data[
        "_v31_calc_started"
    ] = started

    try:

        await _v31_original_do_calculate(
            update,
            context
        )

        context.user_data[
            "last_calculation_seconds"
        ] = round(
            time.perf_counter()
            -
            started,
            4
        )

        v31_mark_dirty(
            context
        )

        await asyncio.to_thread(
            v31_save_project,
            context,
            "calculation"
        )

    finally:

        context.user_data.pop(
            "_v31_calc_started",
            None
        )


# ============================================================
# BUTTON ROUTER
# ============================================================

async def v31_buttons(
    update,
    context
):

    q = update.callback_query

    if q is None:
        return

    data = q.data or ""

    special = {
        "v31_save",
        "v31_projects",
        "v31_copy_project",
        "v31_dashboard",
        "v31_copy_last_member",
        "v31_export_menu",
        "v31_export_xlsx",
        "v31_export_pdf",
        "v31_export_json",
        "v31_export_csv",
    }

    if (
        data.startswith(
            "v31_load_"
        )
        or
        data in special
    ):

        if not await v31_ack_guard(
            q,
            context
        ):
            return

        try:

            handled = (
                await v31_handle_special_callback(
                    update,
                    context,
                    data
                )
            )

            if handled:
                return

        except Exception as exc:

            logger.exception(
                "V3.1 callback failed"
            )

            lang = context.user_data.get(
                "lang",
                "fa"
            )

            await q.edit_message_text(
                TEXT.get(
                    lang,
                    TEXT["fa"]
                )[
                    "calc_error"
                ].format(
                    escape(
                        str(exc)
                    )
                ),
                parse_mode="HTML",
                reply_markup=back_kb(
                    lang
                )
            )

            return

    if data in {
        "project_tools",
        "v31_project_menu"
    }:

        lang = context.user_data.get(
            "lang",
            "fa"
        )

        await q.edit_message_text(
            "🧰 <b>مدیریت پروژه</b>",
            parse_mode="HTML",
            reply_markup=v31_project_keyboard(
                lang
            )
        )

        return

    await _v31_original_buttons(
        update,
        context
    )


# ============================================================
# MAIN MENU UPGRADE
# ============================================================

def v31_main_menu(
    lang
):

    base = _v31_original_main_menu(
        lang
    )

    rows = list(
        base.inline_keyboard
    )

    exists = any(
        any(
            getattr(
                b,
                "callback_data",
                ""
            )
            ==
            "v31_project_menu"
            for b in row
        )
        for row in rows
    )

    if not exists:

        rows.insert(
            max(
                len(rows) - 2,
                0
            ),
            [
                InlineKeyboardButton(
                    "🗂 مدیریت پروژه",
                    callback_data=
                    "v31_project_menu"
                ),

                InlineKeyboardButton(
                    "📊 داشبورد",
                    callback_data=
                    "v31_dashboard"
                ),
            ]
        )

    return InlineKeyboardMarkup(
        rows
    )


# ============================================================
# RESULT MENU UPGRADE
# ============================================================

def v31_result_kb(
    lang
):

    base = _v31_original_result_kb(
        lang
    )

    rows = list(
        base.inline_keyboard
    )

    rows.insert(
        1,
        [
            InlineKeyboardButton(
                "📋 کپی عضو",
                callback_data=
                "v31_copy_last_member"
            ),

            InlineKeyboardButton(
                "💾 ذخیره",
                callback_data=
                "v31_save"
            ),
        ]
    )

    rows.insert(
        2,
        [
            InlineKeyboardButton(
                "📥 Excel",
                callback_data=
                "v31_export_xlsx"
            ),

            InlineKeyboardButton(
                "📄 PDF",
                callback_data=
                "v31_export_pdf"
            ),
        ]
    )

    return InlineKeyboardMarkup(
        rows
    )


# ============================================================
# GLOBAL REPLACEMENTS
# ============================================================

buttons = v31_buttons

do_calculate = v31_do_calculate

main_menu = v31_main_menu

result_kb = v31_result_kb


# ============================================================
# START
# ============================================================

async def v31_start(
    update,
    context
):

    context.user_data.clear()

    context.user_data[
        "lang"
    ] = "fa"

    try:

        context.user_data[
            "_user_id"
        ] = str(
            update.effective_user.id
        )

    except Exception:

        context.user_data[
            "_user_id"
        ] = "anonymous"

    context.user_data[
        "project_id"
    ] = hashlib.sha1(
        (
            f"{v31_user_id(context)}:"
            f"{time.time_ns()}"
        ).encode(
            "utf-8"
        )
    ).hexdigest()[:12]

    context.user_data[
        "project_name"
    ] = "پروژه جدید"

    context.user_data[
        "project_dirty"
    ] = True

    context.user_data[
        "project_results"
    ] = []

    await update.message.reply_text(
        TEXT["fa"]["language"],
        parse_mode="HTML",
        reply_markup=
        language_keyboard()
    )


start = v31_start


# ============================================================
# CLEAN OLD EXPORTS
# ============================================================

def v31_cleanup_old_exports():

    cutoff = (
        time.time()
        -
        V31_EXPORT_TTL_SECONDS
    )

    for root, _dirs, files in os.walk(
        V31_ROOT
    ):

        for name in files:

            if not name.lower().endswith(
                (
                    ".xlsx",
                    ".pdf",
                    ".csv",
                    ".json"
                )
            ):
                continue

            path = (
                Path(root) /
                name
            )

            try:

                if (
                    path.stat().st_mtime
                    <
                    cutoff
                    and
                    path.stat().st_size
                    <=
                    V31_FILE_LIMIT_MB *
                    1024 *
                    1024
                ):

                    if name.lower().endswith(
                        (
                            ".xlsx",
                            ".pdf",
                            ".csv"
                        )
                    ):

                        path.unlink()

            except OSError:
                pass


# ============================================================
# QUICK PRESET CALLBACKS
# ============================================================

_v31_original_buttons_2 = buttons


async def v31_buttons_final(
    update,
    context
):

    q = update.callback_query

    if q is None:
        return

    data = q.data or ""

    lang = context.user_data.get(
        "lang",
        "fa"
    )

    # ========================================================
    # QUICK PRESET
    # ========================================================

    if data.startswith(
        "v31preset_"
    ):

        raw = data[
            len("v31preset_"):
        ]

        try:

            value = parse_number(
                raw
            )

        except Exception:

            await q.answer(
                "مقدار نامعتبر",
                show_alert=True
            )

            return

        # ----------------------------------------------------
        # ROOF
        # ----------------------------------------------------

        if (
            context.user_data.get(
                "kind"
            )
            ==
            "roof"
        ):

            rt = context.user_data.get(
                "roof_type"
            )

            idx = context.user_data.get(
                "roof_step_index"
            )

            if (
                rt is None
                or
                idx is None
                or
                rt not in ROOF_STEPS
            ):

                await q.answer(
                    "مرحله ورودی فعال نیست",
                    show_alert=True
                )

                return

            key, _label, typ = (
                ROOF_STEPS[
                    rt
                ][
                    idx
                ]
            )

            context.user_data.setdefault(
                "roof_values",
                {}
            )[key] = value

            nxt = idx + 1

            if (
                nxt
                <
                len(
                    ROOF_STEPS[
                        rt
                    ]
                )
            ):

                context.user_data[
                    "roof_step_index"
                ] = nxt

                f = ROOF_STEPS[
                    rt
                ][
                    nxt
                ]

                await q.edit_message_text(
                    prompt_text(
                        lang,
                        f[1],
                        nxt + 1,
                        len(
                            ROOF_STEPS[
                                rt
                            ]
                        )
                    ),
                    parse_mode="HTML",
                    reply_markup=
                    step_kb(
                        lang,
                        rt,
                        nxt,
                        True
                    )
                )

            else:

                context.user_data[
                    "roof_step_index"
                ] = None

                await q.edit_message_text(
                    review_roof(
                        context,
                        rt,
                        context.user_data[
                            "roof_values"
                        ],
                        lang
                    ),
                    parse_mode="HTML",
                    reply_markup=
                    review_kb(
                        lang
                    )
                )

        # ----------------------------------------------------
        # NORMAL MEMBER
        # ----------------------------------------------------

        else:

            kind = context.user_data.get(
                "kind"
            )

            idx = context.user_data.get(
                "step_index"
            )

            if (
                kind is None
                or
                idx is None
                or
                kind not in STEPS
            ):

                await q.answer(
                    "مرحله ورودی فعال نیست",
                    show_alert=True
                )

                return

            key, _label, typ = (
                STEPS[
                    kind
                ][
                    idx
                ]
            )

            if (
                value == 0
                and
                typ not in (
                    "diameter_optional",
                    "spacing_optional",
                    "int0"
                )
            ):

                await q.answer(
                    "صفر برای این فیلد مجاز نیست",
                    show_alert=True
                )

                return

            context.user_data.setdefault(
                "values",
                {}
            )[key] = value

            context.user_data.setdefault(
                "history",
                []
            ).append(
                idx
            )

            nxt = idx + 1

            if (
                nxt
                <
                len(
                    STEPS[
                        kind
                    ]
                )
            ):

                context.user_data[
                    "step_index"
                ] = nxt

                f = STEPS[
                    kind
                ][
                    nxt
                ]

                await q.edit_message_text(
                    prompt_text(
                        lang,
                        field_label(
                            f,
                            lang
                        ),
                        nxt + 1,
                        len(
                            STEPS[
                                kind
                            ]
                        )
                    ),
                    parse_mode="HTML",
                    reply_markup=
                    step_kb(
                        lang,
                        kind,
                        nxt
                    )
                )

            else:

                context.user_data[
                    "step_index"
                ] = None

                await q.edit_message_text(
                    review_text_with_context(
                        context,
                        kind,
                        context.user_data[
                            "values"
                        ],
                        lang
                    ),
                    parse_mode="HTML",
                    reply_markup=
                    review_kb(
                        lang
                    )
                )

        v31_mark_dirty(
            context
        )

        return

    await _v31_original_buttons_2(
        update,
        context
    )


buttons = v31_buttons_final


# ============================================================
# CALLBACK ROUTER
# ============================================================

async def v31_callback_router(
    update,
    context
):

    q = update.callback_query

    if q is None:
        return

    try:

        if not await v31_ack_guard(
            q,
            context
        ):

            return

        await q.answer()

    except Exception:
        pass

    try:

        await buttons(
            update,
            context
        )

    except Exception as exc:

        logger.exception(
            "V3.1 callback handling failed: %s",
            exc
        )

        lang = context.user_data.get(
            "lang",
            "fa"
        )

        try:

            await q.edit_message_text(
                TEXT[lang][
                    "calc_error"
                ].format(
                    escape(
                        str(exc)
                    )
                ),
                parse_mode="HTML",
                reply_markup=
                back_kb(
                    lang
                )
            )

        except Exception:

            logger.exception(
                "Could not render callback error"
            )


# ============================================================
# FINAL MAIN
# ============================================================

def main():

    if not BOT_TOKEN:

        raise RuntimeError(
            "BOT_TOKEN environment variable is not set"
        )

    if not RENDER_EXTERNAL_URL:

        raise RuntimeError(
            "RENDER_EXTERNAL_URL environment variable is not set"
        )

    v31_cleanup_old_exports()

    webhook_url = (
        f"{RENDER_EXTERNAL_URL}/telegram"
    )

    app = (
        Application
        .builder()
        .token(
            BOT_TOKEN
        )
        .build()
    )

    app.add_handler(
        CommandHandler(
            "start",
            start
        )
    )

    app.add_handler(
        CallbackQueryHandler(
            v31_callback_router
        )
    )

    app.add_handler(
        MessageHandler(
            filters.TEXT &
            ~filters.COMMAND,
            receive_router
        )
    )

    app.add_error_handler(
        error_handler
    )

    print(
        f"Concrete Structure Quantity Bot - V{V31_VERSION}"
    )

    print(
        f"Port: {PORT}"
    )

    print(
        f"Webhook: {webhook_url}"
    )

    app.run_webhook(
        listen="0.0.0.0",
        port=PORT,
        url_path="telegram",
        webhook_url=webhook_url,
        allowed_updates=
            Update.ALL_TYPES,
        drop_pending_updates=True
    )


# ============================================================
# RUN
# ============================================================

if __name__ == "__main__":
    main()
